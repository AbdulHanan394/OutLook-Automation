from io import BytesIO

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from app.ingest.attachments import attachment_file_type, extract_attachment_text


def test_supported_text_formats_extract_searchable_text():
    assert attachment_file_type("budget.PDF") == "pdf"
    assert attachment_file_type("forecast.xlsx") == "xlsx"
    assert attachment_file_type("invoice.exe") is None

    assert extract_attachment_text("txt", b"Forecast for blue widget: 48 units.") == (
        "Forecast for blue widget: 48 units.",
        "indexed",
    )

    document = Document()
    document.add_paragraph("Quarterly blue widget forecast")
    docx_data = BytesIO()
    document.save(docx_data)
    text, status = extract_attachment_text("docx", docx_data.getvalue())
    assert status == "indexed"
    assert "Quarterly blue widget forecast" in text

    workbook = Workbook()
    workbook.active.append(["Widget", "Forecast"])
    workbook.active.append(["Blue", 48])
    xlsx_data = BytesIO()
    workbook.save(xlsx_data)
    text, status = extract_attachment_text("xlsx", xlsx_data.getvalue())
    assert status == "indexed"
    assert "Blue 48" in text


def test_encrypted_and_invalid_attachments_are_skipped():
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.encrypt("secret")
    pdf_data = BytesIO()
    writer.write(pdf_data)

    assert extract_attachment_text("pdf", pdf_data.getvalue()) == ("", "skipped_encrypted")
    assert extract_attachment_text("docx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1legacy-office") == (
        "",
        "skipped_encrypted",
    )
    assert extract_attachment_text("pdf", b"not a PDF") == ("", "failed_parse")
    assert extract_attachment_text("unsupported", b"content") == ("", "skipped_unsupported")
