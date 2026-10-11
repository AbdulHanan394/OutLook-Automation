"""Bounded, in-memory text extraction for supported Outlook attachment types."""
from io import BytesIO
from itertools import chain
from pathlib import PurePath
from zipfile import BadZipFile, ZipFile

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

SUPPORTED_TYPES = {"pdf", "docx", "xlsx", "txt", "csv"}
MAX_EXTRACTED_CHARS = 100_000
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
MAX_ARCHIVE_MEMBERS = 10_000
OLE_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def attachment_file_type(filename: str) -> str | None:
    suffix = PurePath(filename).suffix.lower().lstrip(".")
    return suffix if suffix in SUPPORTED_TYPES else None


def _archive_status(content: bytes) -> str | None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            members = archive.infolist()
            if any(member.flag_bits & 1 for member in members):
                return "skipped_encrypted"
            uncompressed_size = sum(member.file_size for member in members)
            if (
                len(members) > MAX_ARCHIVE_MEMBERS
                or uncompressed_size > MAX_ARCHIVE_UNCOMPRESSED_BYTES
            ):
                return "skipped_oversized"
    except (BadZipFile, OSError, EOFError):
        return "failed_parse"
    return None


def extract_attachment_text(file_type: str, content: bytes) -> tuple[str, str]:
    if file_type not in SUPPORTED_TYPES:
        return "", "skipped_unsupported"
    try:
        if file_type == "pdf":
            reader = PdfReader(BytesIO(content), strict=False)
            if reader.is_encrypted:
                return "", "skipped_encrypted"
            parts = []
            for page in reader.pages:
                parts.append(page.extract_text() or "")
                if sum(map(len, parts)) >= MAX_EXTRACTED_CHARS:
                    break
            text = "\n".join(parts)
        elif file_type == "docx":
            if content.startswith(OLE_SIGNATURE):
                return "", "skipped_encrypted"
            archive_status = _archive_status(content)
            if archive_status:
                return "", archive_status
            document = Document(BytesIO(content))
            parts = []
            char_count = 0
            paragraph_text = (paragraph.text for paragraph in document.paragraphs)
            cell_text = (
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            )
            for value in chain(paragraph_text, cell_text):
                if char_count >= MAX_EXTRACTED_CHARS:
                    break
                text_part = value[: MAX_EXTRACTED_CHARS - char_count]
                parts.append(text_part)
                char_count += len(text_part)
            text = "\n".join(parts)
        elif file_type == "xlsx":
            if content.startswith(OLE_SIGNATURE):
                return "", "skipped_encrypted"
            archive_status = _archive_status(content)
            if archive_status:
                return "", archive_status
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            try:
                parts = []
                char_count = 0
                for sheet in workbook.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        line = " ".join(str(value) for value in row if value is not None)
                        if line:
                            line = line[: MAX_EXTRACTED_CHARS - char_count]
                            parts.append(line)
                            char_count += len(line)
                        if char_count >= MAX_EXTRACTED_CHARS:
                            break
                    if char_count >= MAX_EXTRACTED_CHARS:
                        break
                text = "\n".join(parts)
            finally:
                workbook.close()
        else:
            text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = content.decode("cp1252")
        except UnicodeDecodeError:
            return "", "failed_parse"
    except Exception:
        return "", "failed_parse"

    text = text[:MAX_EXTRACTED_CHARS].strip()
    if not text:
        return "", "empty"
    return text, "indexed"
